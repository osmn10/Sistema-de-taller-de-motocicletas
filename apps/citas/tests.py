from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from unittest.mock import patch

from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook
from django.utils import timezone

from apps.configuracion.models import HorarioTaller

from apps.productos.models import Producto
from apps.servicios.models import Servicio
from apps.usuarios.models import Mecanico, Usuario
from apps.vehiculos.models import Motocicleta

from .models import Cita, EstadoCita, RepuestoUsado, ServicioCita
from .reportes import calcular_reporte_servicios
from .services import (
    notificar_cita_agendada,
    notificar_cita_cancelada,
    notificar_cita_completada,
    notificar_cita_confirmada,
    notificar_cita_reagendada,
)
from .totales import calcular_detalle_cita
from .ticket_pdf import generar_ticket_pdf


@override_settings(
    EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend',
    TALLER_NOMBRE='Taller de prueba',
    TALLER_DIRECCION='San Miguel, El Salvador',
)
class NotificacionesCitaTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.cliente = Usuario.objects.create_cliente(
            dui='11111111-1',
            password='cliente123',
            nombre='Ana',
            apellido='López',
            telefono='7000-1111',
            email='ana@example.com',
            direccion='San Miguel',
        )
        cls.admin = Usuario.objects.create_admin(
            dui='22222222-2',
            password='admin123',
            nombre='Alex',
            apellido='Admin',
            telefono='7000-2222',
            email='admin@example.com',
        )
        Usuario.objects.create_mecanico(
            dui='33333333-3',
            password='mecanico123',
            nombre='Mario',
            apellido='Mecánico',
            telefono='7000-3333',
            email='mario@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MECANICA_GENERAL,
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-1234',
            cliente=cls.cliente,
            marca='Honda',
            modelo='CG 150',
            anio=2024,
            color='Negro',
        )
        cls.servicio = Servicio.objects.create(
            nombre='Cambio de aceite',
            precio_base='15.00',
            duracion_estimada=30,
        )

    def crear_cita(self, estado='pendiente'):
        cita = Cita.objects.create(
            cliente=self.cliente,
            motocicleta=self.motocicleta,
            fecha=date.today() + timedelta(days=2),
            hora=time(9, 0),
            estado_id=estado,
        )
        ServicioCita.objects.create(
            cita=cita,
            servicio=self.servicio,
            precio_final=self.servicio.precio_base,
        )
        return cita

    def test_correo_de_cita_agendada_incluye_detalles(self):
        cita = self.crear_cita()

        enviado = notificar_cita_agendada(cita.id)

        self.assertTrue(enviado)
        self.assertEqual(len(mail.outbox), 1)
        mensaje = mail.outbox[0]
        self.assertEqual(mensaje.to, ['ana@example.com'])
        self.assertIn(f'Cita #{cita.id} agendada', mensaje.subject)
        contenido_html, tipo_mime = mensaje.alternatives[0]
        self.assertEqual(tipo_mime, 'text/html')
        self.assertIn('Cambio de aceite', contenido_html)
        self.assertIn('M-1234', contenido_html)

    def test_correo_de_cita_confirmada_incluye_direccion(self):
        cita = self.crear_cita(estado='confirmada')

        enviado = notificar_cita_confirmada(cita.id)

        self.assertTrue(enviado)
        mensaje = mail.outbox[0]
        self.assertIn(f'Cita #{cita.id} confirmada', mensaje.subject)
        contenido_html, _ = mensaje.alternatives[0]
        self.assertIn('San Miguel, El Salvador', contenido_html)

    @patch('apps.citas.views.notificar_cita_agendada')
    def test_agendar_dispara_notificacion_despues_de_guardar(self, notificar):
        self.client.force_login(self.cliente)
        fecha = date.today() + timedelta(days=3)
        url = reverse('citas:agendar_cita')
        url = f'{url}?fecha={fecha.isoformat()}&hora=09:00&servicio={self.servicio.id}'

        with self.captureOnCommitCallbacks(execute=True):
            respuesta = self.client.post(
                url,
                {
                    'motocicleta': self.motocicleta.placa,
                    'observaciones': 'Prueba automática',
                },
            )

        self.assertEqual(respuesta.status_code, 302)
        cita = Cita.objects.get(fecha=fecha)
        notificar.assert_called_once_with(cita.id)

    @patch('apps.citas.views.notificar_cita_confirmada')
    def test_confirmar_dispara_segunda_notificacion(self, notificar):
        cita = self.crear_cita()
        self.client.force_login(self.admin)

        with self.captureOnCommitCallbacks(execute=True):
            respuesta = self.client.post(
                reverse('citas:cita_admin_detalle', args=[cita.id]),
                {
                    'accion': 'cambiar_estado',
                    'nuevo_estado': 'confirmada',
                    'motivo': '',
                },
            )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'confirmada')
        notificar.assert_called_once_with(cita.id)

    def test_correo_de_cita_cancelada_incluye_detalles(self):
        cita = self.crear_cita()

        enviado = notificar_cita_cancelada(cita.id)

        self.assertTrue(enviado)
        self.assertEqual(len(mail.outbox), 1)
        mensaje = mail.outbox[0]
        self.assertEqual(mensaje.to, ['ana@example.com'])
        self.assertIn(f'Cita #{cita.id} cancelada', mensaje.subject)
        contenido_html, tipo_mime = mensaje.alternatives[0]
        self.assertEqual(tipo_mime, 'text/html')
        self.assertIn('Cambio de aceite', contenido_html)
        self.assertIn('M-1234', contenido_html)

    def test_correo_de_cita_reagendada_incluye_fecha_anterior(self):
        cita = self.crear_cita(estado='confirmada')
        fecha_anterior = cita.fecha
        hora_anterior = cita.hora
        cita.fecha = cita.fecha + timedelta(days=5)
        cita.hora = time(11, 0)
        cita.save()

        enviado = notificar_cita_reagendada(cita.id, fecha_anterior, hora_anterior)

        self.assertTrue(enviado)
        mensaje = mail.outbox[0]
        self.assertIn(f'Cita #{cita.id} reagendada', mensaje.subject)
        contenido_html, _ = mensaje.alternatives[0]
        self.assertIn(fecha_anterior.strftime('%d/%m/%Y'), contenido_html)
        self.assertIn(cita.fecha.strftime('%d/%m/%Y'), contenido_html)

    @patch('apps.citas.views.notificar_cita_cancelada')
    def test_cancelar_dispara_notificacion_despues_de_guardar(self, notificar):
        cita = self.crear_cita()
        self.client.force_login(self.cliente)

        with self.captureOnCommitCallbacks(execute=True):
            respuesta = self.client.post(
                reverse('citas:cancelar_cita', args=[cita.id]),
            )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'cancelada')
        notificar.assert_called_once_with(cita.id)

    @patch('apps.citas.views.notificar_cita_reagendada')
    def test_reagendar_dispara_notificacion_despues_de_guardar(self, notificar):
        cita = self.crear_cita()
        fecha_anterior = cita.fecha
        hora_anterior = cita.hora
        nueva_fecha = fecha_anterior + timedelta(days=3)
        self.client.force_login(self.cliente)

        with self.captureOnCommitCallbacks(execute=True):
            respuesta = self.client.post(
                reverse('citas:reagendar_cita', args=[cita.id]),
                {'fecha': nueva_fecha.isoformat(), 'hora': '10:00'},
            )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.fecha, nueva_fecha)
        notificar.assert_called_once_with(cita.id, fecha_anterior, hora_anterior)


class FinalizarServicioTests(TestCase):
    """V2SCRUM-24: registrar repuestos usados y finalizar la cita."""

    @classmethod
    def setUpTestData(cls):
        cls.cliente = Usuario.objects.create_cliente(
            dui='44444444-4',
            password='cliente123',
            nombre='Carlos',
            apellido='Pérez',
            telefono='7000-4444',
            email='carlos@example.com',
            direccion='San Salvador',
        )
        cls.admin = Usuario.objects.create_admin(
            dui='55555555-5',
            password='admin123',
            nombre='Alex',
            apellido='Admin',
            telefono='7000-5555',
            email='admin2@example.com',
        )
        cls.mecanico = Usuario.objects.create_mecanico(
            dui='66666666-6',
            password='mecanico123',
            nombre='Mario',
            apellido='Mecánico',
            telefono='7000-6666',
            email='mario2@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MECANICA_GENERAL,
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-9999',
            cliente=cls.cliente,
            marca='Yamaha',
            modelo='FZ 150',
            anio=2023,
            color='Azul',
        )
        cls.servicio = Servicio.objects.create(
            nombre='Ajuste de motor',
            precio_base='40.00',
            duracion_estimada=60,
        )

    def setUp(self):
        self.producto = Producto.objects.create(
            nombre='Pastilla de freno',
            precio='12.50',
            stock_actual=10,
            stock_minimo=2,
        )

    def crear_cita_en_proceso(self):
        cita = Cita.objects.create(
            cliente=self.cliente,
            motocicleta=self.motocicleta,
            mecanico=self.mecanico,
            fecha=date.today(),
            hora=time(9, 0),
            estado_id='en_proceso',
        )
        ServicioCita.objects.create(
            cita=cita,
            servicio=self.servicio,
            precio_final=self.servicio.precio_base,
        )
        return cita

    def test_finalizar_descuenta_inventario_y_completa_la_cita(self):
        cita = self.crear_cita_en_proceso()
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'finalizar_servicio',
                'producto_id': [str(self.producto.id)],
                'cantidad': ['3'],
                'observaciones_cierre': 'Se cambió pastilla de freno trasera.',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.producto.refresh_from_db()

        self.assertEqual(cita.estado_id, 'completada')
        self.assertEqual(cita.observaciones_cierre, 'Se cambió pastilla de freno trasera.')
        self.assertEqual(self.producto.stock_actual, 7)
        self.assertEqual(RepuestoUsado.objects.filter(cita=cita).count(), 1)
        repuesto = RepuestoUsado.objects.get(cita=cita)
        self.assertEqual(repuesto.producto, self.producto)
        self.assertEqual(repuesto.cantidad, 3)

        cambio = cita.cambios_estado.first()
        self.assertEqual(cambio.estado_nuevo_id, 'completada')
        self.assertEqual(cambio.estado_anterior_id, 'en_proceso')

    def test_finalizar_sin_repuestos_es_valido(self):
        cita = self.crear_cita_en_proceso()
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'finalizar_servicio',
                'producto_id': [''],
                'cantidad': [''],
                'observaciones_cierre': 'Servicio de rutina, sin repuestos.',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'completada')
        self.assertEqual(RepuestoUsado.objects.filter(cita=cita).count(), 0)

    def test_finalizar_rechaza_stock_insuficiente(self):
        cita = self.crear_cita_en_proceso()
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'finalizar_servicio',
                'producto_id': [str(self.producto.id)],
                'cantidad': ['999'],
                'observaciones_cierre': '',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.producto.refresh_from_db()

        self.assertEqual(cita.estado_id, 'en_proceso')
        self.assertEqual(self.producto.stock_actual, 10)
        self.assertEqual(RepuestoUsado.objects.filter(cita=cita).count(), 0)

    def test_finalizar_requiere_que_la_cita_este_en_proceso(self):
        cita = self.crear_cita_en_proceso()
        cita.estado_id = 'pendiente'
        cita.save(update_fields=['estado'])
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'finalizar_servicio',
                'producto_id': [''],
                'cantidad': [''],
                'observaciones_cierre': '',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'pendiente')

    def test_cambiar_estado_generico_no_permite_completar_directamente(self):
        cita = self.crear_cita_en_proceso()
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'cambiar_estado',
                'nuevo_estado': 'completada',
                'motivo': '',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'en_proceso')

    def test_finalizar_rechaza_producto_repetido(self):
        cita = self.crear_cita_en_proceso()
        self.client.force_login(self.admin)

        respuesta = self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {
                'accion': 'finalizar_servicio',
                'producto_id': [str(self.producto.id), str(self.producto.id)],
                'cantidad': ['2', '1'],
                'observaciones_cierre': '',
            },
        )

        self.assertEqual(respuesta.status_code, 302)
        cita.refresh_from_db()
        self.producto.refresh_from_db()
        self.assertEqual(cita.estado_id, 'en_proceso')
        self.assertEqual(self.producto.stock_actual, 10)


class TicketYCorreoCierreTests(TestCase):
    """V2SCRUM-30 (ticket PDF) y V2SCRUM-33 (correo de cierre con adjunto)."""

    @classmethod
    def setUpTestData(cls):
        cls.cliente = Usuario.objects.create_cliente(
            dui='11122233-1', password='cliente123', nombre='Laura', apellido='Reyes',
            telefono='7000-1111', email='laura@example.com', direccion='Santa Ana',
        )
        cls.admin = Usuario.objects.create_admin(
            dui='11122233-2', password='admin123', nombre='Ana', apellido='Admin',
            telefono='7000-2222', email='ana-admin@example.com',
        )
        cls.mecanico = Usuario.objects.create_mecanico(
            dui='11122233-3', password='mecanico123', nombre='Beto', apellido='Mecánico',
            telefono='7000-3333', email='beto@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MECANICA_GENERAL,
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-7777', cliente=cls.cliente, marca='Honda', modelo='CB 190',
            anio=2022, color='Rojo',
        )
        cls.servicio = Servicio.objects.create(
            nombre='Cambio de aceite', precio_base='25.00', duracion_estimada=30,
        )

    def setUp(self):
        self.producto = Producto.objects.create(
            nombre='Aceite 10W-40', precio='9.50', stock_actual=15, stock_minimo=2,
        )

    def crear_cita_en_proceso(self):
        cita = Cita.objects.create(
            cliente=self.cliente,
            motocicleta=self.motocicleta,
            mecanico=self.mecanico,
            fecha=date.today(),
            hora=time(10, 0),
            estado_id='en_proceso',
        )
        ServicioCita.objects.create(
            cita=cita, servicio=self.servicio, precio_final=self.servicio.precio_base,
        )
        return cita

    def finalizar(self, cita, **extra):
        datos = {
            'accion': 'finalizar_servicio',
            'producto_id': [str(self.producto.id)],
            'cantidad': ['2'],
            'observaciones_cierre': 'Todo en orden.',
        }
        datos.update(extra)
        self.client.force_login(self.admin)
        return self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]), datos,
        )

    # --- V2SCRUM-30: generación del PDF ---

    def test_generar_ticket_pdf_devuelve_bytes_con_encabezado_pdf(self):
        cita = self.crear_cita_en_proceso()
        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita)
        cita.refresh_from_db()

        detalle = calcular_detalle_cita(cita)
        pdf_bytes = generar_ticket_pdf(cita, detalle)

        self.assertTrue(pdf_bytes.startswith(b'%PDF'))
        self.assertGreater(len(pdf_bytes), 100)

    def test_descargar_ticket_requiere_cita_completada(self):
        cita = self.crear_cita_en_proceso()  # todavía en proceso, no completada
        self.client.force_login(self.cliente)

        respuesta = self.client.get(reverse('citas:descargar_ticket', args=[cita.id]))

        self.assertEqual(respuesta.status_code, 302)

    def test_descargar_ticket_funciona_para_cita_completada(self):
        cita = self.crear_cita_en_proceso()
        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita)
        cita.refresh_from_db()
        self.client.force_login(self.cliente)

        respuesta = self.client.get(reverse('citas:descargar_ticket', args=[cita.id]))

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertIn(f'ticket_cita_{cita.id}.pdf', respuesta['Content-Disposition'])
        self.assertTrue(respuesta.content.startswith(b'%PDF'))

    def test_descargar_ticket_no_permite_ver_cita_de_otro_cliente(self):
        cita = self.crear_cita_en_proceso()
        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita)
        cita.refresh_from_db()

        otro_cliente = Usuario.objects.create_cliente(
            dui='11122233-4', password='x', nombre='Otro', apellido='Cliente',
            telefono='7000-4444', email='otro@example.com', direccion='SS',
        )
        self.client.force_login(otro_cliente)

        respuesta = self.client.get(reverse('citas:descargar_ticket', args=[cita.id]))

        self.assertEqual(respuesta.status_code, 404)

    # --- V2SCRUM-33: correo de cierre ---

    def test_finalizar_dispara_correo_de_cierre_con_pdf_adjunto(self):
        cita = self.crear_cita_en_proceso()

        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita)
        cita.refresh_from_db()

        self.assertEqual(len(mail.outbox), 1)
        mensaje = mail.outbox[0]
        self.assertEqual(mensaje.to, ['laura@example.com'])
        self.assertIn(f'Cita #{cita.id} completada', mensaje.subject)

        self.assertEqual(len(mensaje.attachments), 1)
        nombre_adjunto, contenido, tipo_mime = mensaje.attachments[0]
        self.assertEqual(nombre_adjunto, f'ticket_cita_{cita.id}.pdf')
        self.assertEqual(tipo_mime, 'application/pdf')
        self.assertTrue(contenido.startswith(b'%PDF'))

        contenido_html, _ = mensaje.alternatives[0]
        self.assertIn('Cambio de aceite', contenido_html)
        self.assertIn('Aceite 10W-40', contenido_html)
        from django.template.defaultfilters import floatformat
        total_esperado = floatformat(calcular_detalle_cita(cita)['total_referencia'], 2)
        self.assertIn(total_esperado, contenido_html)

    @patch('apps.citas.views.notificar_cita_completada')
    def test_finalizar_llama_a_notificar_cita_completada_tras_guardar(self, notificar):
        cita = self.crear_cita_en_proceso()

        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita)

        notificar.assert_called_once_with(cita.id)

    def test_notificar_cita_completada_devuelve_true(self):
        cita = self.crear_cita_en_proceso()
        with self.captureOnCommitCallbacks(execute=True):
            self.finalizar(cita, producto_id=[''], cantidad=[''])
        cita.refresh_from_db()
        mail.outbox.clear()

        resultado = notificar_cita_completada(cita.id)

        self.assertTrue(resultado)
        self.assertEqual(len(mail.outbox), 1)


class EstadosCitaTests(TestCase):
    """Catálogo configurable de estados de cita: CRUD del admin y flujo automático por tipo."""

    @classmethod
    def setUpTestData(cls):
        cls.cliente = Usuario.objects.create_cliente(
            dui='11111111-1',
            password='cliente123',
            nombre='Ana',
            apellido='López',
            telefono='7000-1111',
            email='ana@example.com',
            direccion='San Miguel',
        )
        cls.admin = Usuario.objects.create_admin(
            dui='22222222-2',
            password='admin123',
            nombre='Alex',
            apellido='Admin',
            telefono='7000-2222',
            email='admin@example.com',
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-1234',
            cliente=cls.cliente,
            marca='Honda',
            modelo='CG 150',
            anio=2024,
            color='Negro',
        )

    def setUp(self):
        self.client.force_login(self.admin)
        self.url = reverse('citas:estados_cita')

    def crear_cita(self, estado):
        return Cita.objects.create(
            cliente=self.cliente,
            motocicleta=self.motocicleta,
            fecha=date.today() + timedelta(days=2),
            hora=time(9, 0),
            estado_id=estado,
        )

    def crear_estado_proceso(self, nombre='En cambio de aceite'):
        self.client.post(self.url, {'accion': 'guardar', 'nombre': nombre, 'tipo': 'proceso'})
        return EstadoCita.objects.get(nombre=nombre)

    def cambiar_estado(self, cita, nuevo_estado, motivo=''):
        return self.client.post(
            reverse('citas:cita_admin_detalle', args=[cita.id]),
            {'accion': 'cambiar_estado', 'nuevo_estado': nuevo_estado, 'motivo': motivo},
        )

    def siguientes(self, codigo):
        return list(EstadoCita.objects.get(codigo=codigo).siguientes_posibles().values_list('codigo', flat=True))

    # --- migración y flujo ---

    def test_migracion_carga_los_cinco_estados_existentes(self):
        tipos = dict(EstadoCita.objects.values_list('codigo', 'tipo'))
        self.assertEqual(tipos, {
            'pendiente': 'inicio',
            'confirmada': 'proceso',
            'en_proceso': 'proceso',
            'completada': 'completado',
            'cancelada': 'cancelado',
        })

    def test_cita_nueva_recibe_el_estado_de_inicio(self):
        cita = Cita.objects.create(
            cliente=self.cliente, motocicleta=self.motocicleta,
            fecha=date.today() + timedelta(days=2), hora=time(9, 0),
        )
        self.assertEqual(cita.estado_id, 'pendiente')

    def test_flujo_con_los_estados_originales(self):
        self.assertEqual(self.siguientes('pendiente'), ['confirmada', 'cancelada'])
        self.assertEqual(self.siguientes('confirmada'), ['en_proceso', 'cancelada'])
        self.assertEqual(self.siguientes('en_proceso'), ['completada', 'cancelada'])
        self.assertEqual(self.siguientes('completada'), [])
        self.assertEqual(self.siguientes('cancelada'), [])

    def test_estado_de_proceso_nuevo_entra_al_flujo(self):
        nuevo = self.crear_estado_proceso()

        self.assertEqual(self.siguientes('confirmada'), ['en_proceso', nuevo.codigo, 'cancelada'])
        self.assertEqual(self.siguientes('en_proceso'), [nuevo.codigo, 'completada', 'cancelada'])
        self.assertEqual(self.siguientes(nuevo.codigo), ['en_proceso', 'completada', 'cancelada'])
        # desde pendiente se sigue pasando solo a confirmada
        self.assertEqual(self.siguientes('pendiente'), ['confirmada', 'cancelada'])

    def test_estado_desactivado_sale_del_flujo(self):
        nuevo = self.crear_estado_proceso()
        nuevo.activo = False
        nuevo.save()

        self.assertNotIn(nuevo.codigo, self.siguientes('confirmada'))
        self.assertNotIn(nuevo.codigo, self.siguientes('en_proceso'))

    # --- cambio de estado en la cita ---

    def test_rechaza_transicion_que_no_permite_el_flujo(self):
        cita = self.crear_cita('pendiente')

        self.cambiar_estado(cita, 'en_proceso')

        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'pendiente')
        self.assertFalse(cita.cambios_estado.exists())

    def test_permite_pasar_a_un_estado_de_proceso_nuevo(self):
        nuevo = self.crear_estado_proceso()
        cita = self.crear_cita('confirmada')

        self.cambiar_estado(cita, nuevo.codigo)

        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, nuevo.codigo)
        cambio = cita.cambios_estado.first()
        self.assertEqual(cambio.estado_anterior_id, 'confirmada')
        self.assertEqual(cambio.estado_nuevo_id, nuevo.codigo)

    def test_no_permite_pasar_a_un_estado_desactivado(self):
        nuevo = self.crear_estado_proceso()
        nuevo.activo = False
        nuevo.save()
        cita = self.crear_cita('confirmada')

        self.cambiar_estado(cita, nuevo.codigo)

        cita.refresh_from_db()
        self.assertEqual(cita.estado_id, 'confirmada')

    def test_confirmar_envia_correo_y_pasar_entre_procesos_no(self):
        nuevo = self.crear_estado_proceso()

        with patch('apps.citas.views.notificar_cita_confirmada') as notificar:
            cita = self.crear_cita('pendiente')
            with self.captureOnCommitCallbacks(execute=True):
                self.cambiar_estado(cita, 'confirmada')
            notificar.assert_called_once_with(cita.id)

            notificar.reset_mock()
            with self.captureOnCommitCallbacks(execute=True):
                self.cambiar_estado(cita, nuevo.codigo)
            notificar.assert_not_called()

    def test_cliente_no_puede_cancelar_en_estado_de_proceso_nuevo(self):
        nuevo = self.crear_estado_proceso()
        cita = self.crear_cita(nuevo.codigo)

        self.assertFalse(cita.puede_cancelarse())

    # --- CRUD del catálogo ---

    def test_crear_estado_de_proceso(self):
        nuevo = self.crear_estado_proceso('En calibración')

        self.assertEqual(nuevo.codigo, 'en_calibracion')
        self.assertEqual(nuevo.tipo, EstadoCita.TIPO_PROCESO)
        self.assertEqual(nuevo.color, EstadoCita.COLORES[EstadoCita.TIPO_PROCESO])
        self.assertEqual(nuevo.orden, 6)
        self.assertTrue(nuevo.activo)
        self.assertFalse(nuevo.permite_cancelar)

    def test_crear_requiere_nombre(self):
        self.client.post(self.url, {'accion': 'guardar', 'nombre': '', 'tipo': 'proceso'})

        self.assertEqual(EstadoCita.objects.count(), 5)

    def test_crear_no_permite_nombre_repetido(self):
        self.client.post(self.url, {'accion': 'guardar', 'nombre': 'en proceso', 'tipo': 'proceso'})

        self.assertEqual(EstadoCita.objects.count(), 5)

    def test_no_permite_un_segundo_estado_de_tipo_unico(self):
        for tipo in ['inicio', 'completado', 'cancelado']:
            self.client.post(self.url, {'accion': 'guardar', 'nombre': f'Otro {tipo}', 'tipo': tipo})

        self.assertEqual(EstadoCita.objects.count(), 5)

    def test_editar_cambia_el_nombre_y_conserva_el_codigo(self):
        self.client.post(self.url, {
            'accion': 'guardar', 'codigo': 'en_proceso', 'nombre': 'En reparación', 'tipo': 'proceso',
        })

        estado = EstadoCita.objects.get(codigo='en_proceso')
        self.assertEqual(estado.nombre, 'En reparación')

    def test_no_permite_cambiar_el_tipo_de_un_estado_unico(self):
        self.client.post(self.url, {
            'accion': 'guardar', 'codigo': 'completada', 'nombre': 'Completada', 'tipo': 'proceso',
        })

        self.assertEqual(EstadoCita.objects.get(codigo='completada').tipo, EstadoCita.TIPO_COMPLETADO)

    def test_no_permite_desactivar_un_estado_unico(self):
        self.client.post(self.url, {'accion': 'desactivar', 'codigo': 'cancelada'})

        self.assertTrue(EstadoCita.objects.get(codigo='cancelada').activo)

    def test_no_permite_desactivar_un_estado_con_citas(self):
        self.crear_cita('en_proceso')

        self.client.post(self.url, {'accion': 'desactivar', 'codigo': 'en_proceso'})

        self.assertTrue(EstadoCita.objects.get(codigo='en_proceso').activo)

    def test_desactivar_y_activar_estado_sin_citas(self):
        nuevo = self.crear_estado_proceso()

        self.client.post(self.url, {'accion': 'desactivar', 'codigo': nuevo.codigo})
        nuevo.refresh_from_db()
        self.assertFalse(nuevo.activo)

        self.client.post(self.url, {'accion': 'activar', 'codigo': nuevo.codigo})
        nuevo.refresh_from_db()
        self.assertTrue(nuevo.activo)

    def test_solo_el_admin_gestiona_estados(self):
        self.client.force_login(self.cliente)

        respuesta = self.client.post(self.url, {'accion': 'guardar', 'nombre': 'Intruso', 'tipo': 'proceso'})

        self.assertRedirects(respuesta, reverse('core:home'), fetch_redirect_response=False)
        self.assertFalse(EstadoCita.objects.filter(nombre='Intruso').exists())


class HorariosPasadosTests(TestCase):
    """No se puede agendar ni reagendar en un horario de hoy que ya empezó."""

    # lunes fijo para no depender del día en que corran los tests
    HOY = date(2026, 10, 5)

    @classmethod
    def setUpTestData(cls):
        cls.cliente = Usuario.objects.create_cliente(
            dui='11111111-1',
            password='cliente123',
            nombre='Ana',
            apellido='López',
            telefono='7000-1111',
            email='ana@example.com',
            direccion='San Miguel',
        )
        Usuario.objects.create_mecanico(
            dui='33333333-3',
            password='mecanico123',
            nombre='Mario',
            apellido='Mecánico',
            telefono='7000-3333',
            email='mario@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MECANICA_GENERAL,
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-1234',
            cliente=cls.cliente,
            marca='Honda',
            modelo='CG 150',
            anio=2024,
            color='Negro',
        )
        cls.servicio = Servicio.objects.create(
            nombre='Cambio de aceite',
            precio_base='15.00',
            duracion_estimada=30,
        )
        for dia in [cls.HOY, cls.HOY + timedelta(days=1)]:
            HorarioTaller.objects.update_or_create(
                dia_semana=dia.weekday(),
                defaults={'abierto': True, 'hora_apertura': time(8, 0), 'hora_cierre': time(17, 0)},
            )

    def setUp(self):
        self.client.force_login(self.cliente)

    def a_las(self, hora, minuto):
        """Congela el reloj del sistema en HOY a la hora indicada (hora de El Salvador)."""
        momento = timezone.make_aware(datetime.combine(self.HOY, time(hora, minuto)))
        return patch('django.utils.timezone.now', return_value=momento)

    def consultar_disponibilidad(self):
        return self.client.get(reverse('citas:disponibilidad'), {
            'fecha': self.HOY.isoformat(),
            'servicio': self.servicio.id,
        })

    def agendar(self, hora):
        url = reverse('citas:agendar_cita')
        return self.client.post(
            f'{url}?fecha={self.HOY.isoformat()}&hora={hora}&servicio={self.servicio.id}',
            {'motocicleta': self.motocicleta.placa},
        )

    def test_disponibilidad_de_hoy_empieza_en_el_siguiente_horario(self):
        with self.a_las(9, 30):
            respuesta = self.consultar_disponibilidad()

        horas = [slot['hora'] for slot in respuesta.context['slots']]
        self.assertEqual(horas[0], time(10, 0))
        self.assertNotIn(time(9, 30), horas)
        self.assertNotIn(time(8, 0), horas)

    def test_disponibilidad_de_hoy_sin_horarios_restantes(self):
        with self.a_las(22, 41):
            respuesta = self.consultar_disponibilidad()

        self.assertEqual(respuesta.context['slots'], [])
        self.assertIn('Ya no hay horarios disponibles para hoy', respuesta.context['mensaje'])

    def test_disponibilidad_de_otro_dia_muestra_todos_los_horarios(self):
        with self.a_las(22, 41):
            respuesta = self.client.get(reverse('citas:disponibilidad'), {
                'fecha': (self.HOY + timedelta(days=1)).isoformat(),
                'servicio': self.servicio.id,
            })

        self.assertEqual(respuesta.context['slots'][0]['hora'], time(8, 0))

    def test_no_permite_agendar_en_horario_de_hoy_que_ya_paso(self):
        with self.a_las(9, 30):
            respuesta = self.agendar('07:00')

        self.assertRedirects(respuesta, reverse('citas:disponibilidad'), fetch_redirect_response=False)
        self.assertFalse(Cita.objects.exists())

    def test_no_permite_agendar_en_el_horario_que_esta_empezando(self):
        with self.a_las(9, 30):
            self.agendar('09:30')

        self.assertFalse(Cita.objects.exists())

    def test_permite_agendar_en_horario_futuro_de_hoy(self):
        with self.a_las(9, 30):
            self.agendar('10:00')

        cita = Cita.objects.get()
        self.assertEqual(cita.fecha, self.HOY)
        self.assertEqual(cita.hora, time(10, 0))

    def test_no_permite_reagendar_a_horario_de_hoy_que_ya_paso(self):
        cita = Cita.objects.create(
            cliente=self.cliente,
            motocicleta=self.motocicleta,
            fecha=self.HOY + timedelta(days=2),
            hora=time(9, 0),
        )
        ServicioCita.objects.create(cita=cita, servicio=self.servicio, precio_final=self.servicio.precio_base)

        with self.a_las(9, 30):
            self.client.post(
                reverse('citas:reagendar_cita', args=[cita.id]),
                {'fecha': self.HOY.isoformat(), 'hora': '08:00'},
            )

        cita.refresh_from_db()
        self.assertEqual(cita.fecha, self.HOY + timedelta(days=2))
        self.assertEqual(cita.hora, time(9, 0))


class ReporteServiciosTests(TestCase):
    """V2SCRUM-39: reporte de servicios por período, filtros y export."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = Usuario.objects.create_admin(
            dui='20000000-1', password='admin123', nombre='Ana', apellido='Admin',
            telefono='7000-0001', email='ana-admin@example.com',
        )
        cls.cliente = Usuario.objects.create_cliente(
            dui='20000000-2', password='cliente123', nombre='Pedro', apellido='Cliente',
            telefono='7000-0002', email='pedro@example.com', direccion='San Miguel',
        )
        cls.mecanico1 = Usuario.objects.create_mecanico(
            dui='20000000-3', password='mecanico123', nombre='Beto', apellido='Alfa',
            telefono='7000-0003', email='beto@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MECANICA_GENERAL,
        )
        cls.mecanico2 = Usuario.objects.create_mecanico(
            dui='20000000-4', password='mecanico123', nombre='Carla', apellido='Beta',
            telefono='7000-0004', email='carla@example.com',
            especialidad=Mecanico.ESPECIALIDAD_MOTOR,
        )
        cls.motocicleta = Motocicleta.objects.create(
            placa='M-9000', cliente=cls.cliente, marca='Honda', modelo='CG 150',
            anio=2023, color='Negro',
        )
        cls.aceite = Servicio.objects.create(
            nombre='Cambio de aceite', precio_base='15.00', duracion_estimada=30,
        )
        cls.afinamiento = Servicio.objects.create(
            nombre='Afinamiento general', precio_base='20.00', duracion_estimada=60,
        )

        est_inicio = EstadoCita.objects.get(tipo=EstadoCita.TIPO_INICIO)
        est_completado = EstadoCita.objects.get(tipo=EstadoCita.TIPO_COMPLETADO)
        est_cancelado = EstadoCita.objects.get(tipo=EstadoCita.TIPO_CANCELADO)
        cls.est_inicio, cls.est_completado, cls.est_cancelado = est_inicio, est_completado, est_cancelado

        hoy = date.today()

        # Dentro del período, completadas: 2 servicios de aceite + 1 de afinamiento,
        # repartidas 2 citas con mecanico1.
        cls._crear_cita(hoy, cls.est_completado, cls.mecanico1, [cls.aceite])
        cls._crear_cita(hoy, cls.est_completado, cls.mecanico1, [cls.aceite, cls.afinamiento])
        # Dentro del período, pendiente (cuenta como demanda, no como ingreso/completada).
        cls._crear_cita(hoy, cls.est_inicio, cls.mecanico2, [cls.aceite])
        # Dentro del período, cancelada: no debe contar para nada.
        cls._crear_cita(hoy, cls.est_cancelado, cls.mecanico2, [cls.aceite])
        # Fuera del período (60 días atrás): no debe contar para nada.
        cls._crear_cita(hoy - timedelta(days=60), cls.est_completado, cls.mecanico1, [cls.aceite])

    @classmethod
    def _crear_cita(cls, fecha, estado, mecanico, servicios):
        cita = Cita.objects.create(
            cliente=cls.cliente, motocicleta=cls.motocicleta, mecanico=mecanico,
            fecha=fecha, hora=time(9, 0), estado=estado,
        )
        for servicio in servicios:
            ServicioCita.objects.create(
                cita=cita, servicio=servicio, precio_final=servicio.precio_base,
            )
        return cita

    # --- cálculo (apps/citas/reportes.py) ---

    def test_kpis_del_periodo_sin_filtros(self):
        datos = calcular_reporte_servicios(date.today(), date.today())

        self.assertEqual(datos['total_citas'], 3)  # completadas + pendiente, sin cancelada
        self.assertEqual(datos['servicios_completados'], 2)
        self.assertEqual(datos['ingresos_totales'], Decimal('50.00'))  # 15 + (15+20)

    def test_servicios_mas_solicitados_ordenados_por_cantidad(self):
        datos = calcular_reporte_servicios(date.today(), date.today())
        top = datos['servicios_top']

        self.assertEqual(top[0]['servicio__nombre'], 'Cambio de aceite')
        self.assertEqual(top[0]['cantidad'], 3)  # 2 completadas + 1 pendiente
        self.assertEqual(top[1]['servicio__nombre'], 'Afinamiento general')
        self.assertEqual(top[1]['cantidad'], 1)

    def test_mecanico_con_mas_servicios(self):
        datos = calcular_reporte_servicios(date.today(), date.today())
        top = datos['mecanicos_top']

        self.assertEqual(top[0]['mecanico__dui'], self.mecanico1.dui)
        self.assertEqual(top[0]['cantidad'], 2)
        self.assertEqual(top[1]['mecanico__dui'], self.mecanico2.dui)
        self.assertEqual(top[1]['cantidad'], 1)

    def test_cancelada_no_cuenta_para_nada(self):
        datos = calcular_reporte_servicios(date.today(), date.today())
        total_aceite = sum(
            f['cantidad'] for f in datos['servicios_top'] if f['servicio__nombre'] == 'Cambio de aceite'
        )
        # 3 (completadas + pendiente), NO 4: la cancelada quedó fuera.
        self.assertEqual(total_aceite, 3)

    def test_fuera_de_rango_no_cuenta(self):
        datos = calcular_reporte_servicios(date.today(), date.today())
        self.assertEqual(datos['total_citas'], 3)  # no las 4 activas totales

    def test_filtro_por_servicio_acota_las_citas(self):
        datos = calcular_reporte_servicios(
            date.today(), date.today(), servicio_id=self.afinamiento.id,
        )
        # Solo la cita que tiene afinamiento (la que además tiene aceite).
        self.assertEqual(datos['total_citas'], 1)
        self.assertEqual(datos['ingresos_totales'], Decimal('35.00'))

    def test_filtro_por_mecanico_acota_las_citas(self):
        datos = calcular_reporte_servicios(
            date.today(), date.today(), mecanico_dui=self.mecanico2.dui,
        )
        self.assertEqual(datos['total_citas'], 1)  # solo la pendiente (la cancelada no cuenta)
        self.assertEqual(datos['servicios_completados'], 0)
        self.assertEqual(datos['ingresos_totales'], Decimal('0.00'))

    # --- vista y permisos ---

    def test_requiere_admin(self):
        self.client.force_login(self.cliente)
        respuesta = self.client.get(reverse('citas:reporte_servicios'))
        self.assertEqual(respuesta.status_code, 302)

    def test_admin_ve_el_reporte(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('citas:reporte_servicios'))
        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta.context['total_citas'], 3)

    # --- exports ---

    def test_export_pdf(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('citas:reporte_servicios'), {'export': 'pdf'})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(respuesta['Content-Type'], 'application/pdf')
        self.assertIn('attachment', respuesta['Content-Disposition'])
        self.assertTrue(respuesta.content.startswith(b'%PDF'))

    def test_export_excel_es_un_xlsx_valido_con_los_datos(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('citas:reporte_servicios'), {'export': 'xlsx'})

        self.assertEqual(respuesta.status_code, 200)
        self.assertEqual(
            respuesta['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(
            libro.sheetnames, ['Resumen', 'Servicios más solicitados', 'Mecánicos'],
        )
        resumen = libro['Resumen']
        self.assertEqual(resumen['A4'].value, 'Citas del período')
        self.assertEqual(resumen['B4'].value, 3)
        self.assertEqual(resumen['B6'].value, 50.0)

    def test_export_respeta_los_filtros_aplicados(self):
        self.client.force_login(self.admin)
        respuesta = self.client.get(reverse('citas:reporte_servicios'), {
            'export': 'xlsx', 'mecanico': self.mecanico2.dui,
        })

        libro = load_workbook(BytesIO(respuesta.content))
        self.assertEqual(libro['Resumen']['B4'].value, 1)
