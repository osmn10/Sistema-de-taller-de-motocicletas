"""V2SCRUM-39: genera el Excel del reporte de servicios por período.

Mismo criterio que ``reporte_pdf.py``: se arma "al vuelo" a partir del dict
de ``calcular_reporte_servicios``, sin persistir nada en disco/BD.
"""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font


def generar_reporte_excel(desde, hasta, datos) -> bytes:
    """Arma el .xlsx del reporte. ``datos`` es el resultado de
    ``calcular_reporte_servicios``."""

    libro = Workbook()

    resumen = libro.active
    resumen.title = 'Resumen'
    resumen.append(['Reporte de servicios'])
    resumen['A1'].font = Font(bold=True, size=14)
    resumen.append([f'Período: {desde.strftime("%d/%m/%Y")} — {hasta.strftime("%d/%m/%Y")}'])
    resumen.append([])
    resumen.append(['Citas del período', datos['total_citas']])
    resumen.append(['Servicios completados', datos['servicios_completados']])
    resumen.append(['Ingresos totales', float(datos['ingresos_totales'])])
    resumen.column_dimensions['A'].width = 28
    resumen.column_dimensions['B'].width = 16

    hoja_servicios = libro.create_sheet('Servicios más solicitados')
    hoja_servicios.append(['Servicio', 'Cantidad'])
    for celda in hoja_servicios[1]:
        celda.font = Font(bold=True)
    for fila in datos['servicios_top']:
        hoja_servicios.append([fila['servicio__nombre'], fila['cantidad']])
    hoja_servicios.column_dimensions['A'].width = 35
    hoja_servicios.column_dimensions['B'].width = 12

    hoja_mecanicos = libro.create_sheet('Mecánicos')
    hoja_mecanicos.append(['Mecánico', 'Cantidad'])
    for celda in hoja_mecanicos[1]:
        celda.font = Font(bold=True)
    for fila in datos['mecanicos_top']:
        nombre = f"{fila['mecanico__nombre']} {fila['mecanico__apellido']}"
        hoja_mecanicos.append([nombre, fila['cantidad']])
    hoja_mecanicos.column_dimensions['A'].width = 30
    hoja_mecanicos.column_dimensions['B'].width = 12

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
