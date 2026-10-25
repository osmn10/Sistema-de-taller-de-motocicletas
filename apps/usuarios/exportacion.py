"""V2SCRUM-41: genera el Excel con el listado de clientes y sus vehículos.

Igual que los demás reportes del sistema (ticket PDF, reporte de servicios), el
archivo se arma "al vuelo" a partir de la base de datos y no se guarda en disco.
"""

import unicodedata
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Cliente

ENCABEZADOS = [
    'DUI', 'Nombre', 'Apellido', 'Teléfono', 'Correo', 'Dirección', 'Estado del cliente',
    'Placa', 'Marca', 'Modelo', 'Año', 'Color', 'Estado del vehículo',
]
ANCHOS = [14, 20, 20, 13, 32, 38, 18, 10, 16, 18, 8, 14, 20]


def _clave_orden(texto):
    """Orden alfabético sin distinguir tildes ni mayúsculas (Álvarez junto a Alvarez),
    igual en SQLite y en PostgreSQL."""
    base = unicodedata.normalize('NFD', texto or '')
    return ''.join(c for c in base if not unicodedata.combining(c)).casefold()


def _texto(valor):
    """Devuelve el valor como texto. Se fuerza a texto en la celda para que un
    dato que empiece con "=" nunca se interprete como fórmula de Excel."""
    return '' if valor is None else str(valor)


def generar_excel_clientes_vehiculos() -> bytes:
    """Arma el .xlsx con una fila por vehículo; un cliente sin vehículos
    aparece en una sola fila con los datos del vehículo vacíos."""

    libro = Workbook()
    hoja = libro.active
    hoja.title = 'Clientes y vehículos'

    hoja.append(ENCABEZADOS)
    for celda in hoja[1]:
        celda.font = Font(bold=True)
        celda.fill = PatternFill('solid', fgColor='BDD6EE')
        celda.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    clientes = sorted(
        Cliente.objects.prefetch_related('motocicletas'),
        key=lambda c: (_clave_orden(c.apellido), _clave_orden(c.nombre), c.dui),
    )
    for cliente in clientes:
        datos_cliente = [
            cliente.dui, cliente.nombre, cliente.apellido, cliente.telefono,
            cliente.email, cliente.direccion, 'Activo' if cliente.activo else 'Inactivo',
        ]
        motos = sorted(cliente.motocicletas.all(), key=lambda m: m.placa)
        if not motos:
            hoja.append([_texto(v) for v in datos_cliente] + [''] * 6)
            continue
        for moto in motos:
            datos_moto = [
                moto.placa, moto.marca, moto.modelo, moto.anio, moto.color,
                'Activo' if moto.activo else 'Inactivo',
            ]
            hoja.append([_texto(v) for v in datos_cliente] + datos_moto)

    # Todas las celdas de texto quedan como texto (no como fórmula).
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, str):
                celda.data_type = 's'

    for indice, ancho in enumerate(ANCHOS, start=1):
        hoja.column_dimensions[hoja.cell(1, indice).column_letter].width = ancho
    hoja.freeze_panes = 'A2'

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()