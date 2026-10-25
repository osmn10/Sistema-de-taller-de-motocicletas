"""Tests de la app usuarios.

Cubren la política de contraseña segura y el flujo de recuperación de
contraseña por correo (V2SCRUM-34).
"""

from io import BytesIO

from django.contrib.auth.tokens import default_token_generator
from django.core import mail
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils.encoding import force_bytes
from django.utils.http import urlsafe_base64_encode
from openpyxl import load_workbook

from apps.usuarios.models import Cliente, Usuario
from apps.vehiculos.models import Motocicleta
from apps.usuarios.validators import validar_password_segura


class PoliticaPasswordTests(TestCase):
    """La política exige longitud, mayúscula, minúscula, número y símbolo."""

    def test_password_valida_no_devuelve_errores(self):
        self.assertEqual(validar_password_segura('Abcdef1$'), [])

    def test_password_corta_falla(self):
        self.assertTrue(any('8 caracteres' in e for e in validar_password_segura('Ab1$')))

    def test_password_sin_mayuscula_falla(self):
        self.assertTrue(any('mayúscula' in e for e in validar_password_segura('abcdefg1$')))

    def test_password_sin_minuscula_falla(self):
        self.assertTrue(any('minúscula' in e for e in validar_password_segura('ABCDEFG1$')))

    def test_password_sin_numero_falla(self):
        self.assertTrue(any('número' in e for e in validar_password_segura('Abcdefg$')))

    def test_password_sin_simbolo_falla(self):
        self.assertTrue(any('símbolo' in e for e in validar_password_segura('Abcdefg1')))


@override_settings(EMAIL_BACKEND='django.core.mail.backends.locmem.EmailBackend')
class RecuperacionPasswordTests(TestCase):
    """Flujo completo de recuperación por correo."""

    def setUp(self):
        self.cliente = Cliente(
            dui='12345678-9', nombre='Ana', apellido='Reyes',
            telefono='7777-7777', email='ana@example.com', direccion='San Salvador',
        )
        self.cliente.set_password('ClaveVieja1$')
        self.cliente.save()

    def _enlace_valido(self):
        uid = urlsafe_base64_encode(force_bytes(self.cliente.pk))
        token = default_token_generator.make_token(self.cliente)
        return reverse(
            'usuarios:password_reset_confirmar',
            kwargs={'uidb64': uid, 'token': token},
        )

    def test_solicitud_con_correo_existente_envia_email(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': 'ana@example.com'},
        )
        self.assertRedirects(resp, reverse('usuarios:login'))
        self.assertEqual(len(mail.outbox), 1)
        self.assertIn('ana@example.com', mail.outbox[0].to)

    def test_solicitud_con_dui_existente_envia_email(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': '12345678-9'},
        )
        self.assertRedirects(resp, reverse('usuarios:login'))
        self.assertEqual(len(mail.outbox), 1)

    def test_solicitud_con_cuenta_inexistente_no_envia_ni_revela(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': 'nadie@example.com'},
        )
        self.assertRedirects(resp, reverse('usuarios:login'))
        self.assertEqual(len(mail.outbox), 0)

    def test_dui_bien_formado_inexistente_sigue_generico(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': '99999999-9'},
        )
        self.assertRedirects(resp, reverse('usuarios:login'))
        self.assertEqual(len(mail.outbox), 0)

    def test_dui_mal_formado_muestra_error_de_formato(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': '123'},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'formato')
        self.assertEqual(len(mail.outbox), 0)

    def test_correo_mal_formado_muestra_error_de_formato(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': 'ana@correo'},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, 'formato')
        self.assertEqual(len(mail.outbox), 0)

    def test_campo_vacio_muestra_error(self):
        resp = self.client.post(
            reverse('usuarios:password_reset_solicitar'),
            {'identificador': ''},
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(mail.outbox), 0)

    def test_confirmar_cambia_password_y_notifica(self):
        resp = self.client.post(
            self._enlace_valido(),
            {'contrasena': 'ClaveNueva9#', 'confirmar': 'ClaveNueva9#'},
        )
        self.assertRedirects(resp, reverse('usuarios:login'))
        self.cliente.refresh_from_db()
        self.assertTrue(self.cliente.check_password('ClaveNueva9#'))
        self.assertEqual(len(mail.outbox), 1)

    def test_confirmar_rechaza_password_insegura(self):
        resp = self.client.post(
            self._enlace_valido(),
            {'contrasena': 'insegura', 'confirmar': 'insegura'},
        )
        self.assertEqual(resp.status_code, 200)
        self.cliente.refresh_from_db()
        self.assertTrue(self.cliente.check_password('ClaveVieja1$'))

    def test_confirmar_rechaza_si_no_coinciden(self):
        resp = self.client.post(
            self._enlace_valido(),
            {'contrasena': 'ClaveNueva9#', 'confirmar': 'Otra9#Cosa'},
        )
        self.assertEqual(resp.status_code, 200)
        self.cliente.refresh_from_db()
        self.assertTrue(self.cliente.check_password('ClaveVieja1$'))

    def test_token_invalido_redirige_a_solicitar(self):
        url = reverse(
            'usuarios:password_reset_confirmar',
            kwargs={'uidb64': 'xxx', 'token': 'token-invalido'},
        )
        resp = self.client.get(url)
        self.assertRedirects(resp, reverse('usuarios:password_reset_solicitar'))

    def test_token_no_se_puede_reutilizar(self):
        url = self._enlace_valido()
        self.client.post(url, {'contrasena': 'ClaveNueva9#', 'confirmar': 'ClaveNueva9#'})
        resp = self.client.get(url)
        self.assertRedirects(resp, reverse('usuarios:password_reset_solicitar'))


class ExportarClientesVehiculosTests(TestCase):
    """V2SCRUM-41: exportación a Excel del listado de clientes y vehículos."""

    @classmethod
    def setUpTestData(cls):
        cls.admin = Usuario.objects.create_admin(
            dui='00000001-1', password='Admin123$', nombre='Ada', apellido='Admin',
            telefono='7000-0001', email='admin@example.com',
        )
        cls.mecanico = Usuario.objects.create_mecanico(
            dui='00000002-2', password='Mecan123$', especialidad='motor', nombre='Mario',
            apellido='Mecánico', telefono='7000-0002', email='mecanico@example.com',
        )
        cls.ana = Cliente(
            dui='11111111-1', nombre='Ana', apellido='Álvarez', telefono='7111-1111',
            email='ana@example.com', direccion='San Salvador',
        )
        cls.ana.set_password('Cliente123$')
        cls.ana.save()
        cls.beto = Cliente(
            dui='22222222-2', nombre='Beto', apellido='Núñez', telefono='7222-2222',
            email='beto@example.com', direccion='Santa Ana', activo=False,
        )
        cls.beto.set_password('Cliente123$')
        cls.beto.save()
        cls.moto1 = Motocicleta.objects.create(
            placa='M-1111', cliente=cls.ana, marca='Honda', modelo='CB190', anio=2020, color='Rojo',
        )
        cls.moto2 = Motocicleta.objects.create(
            placa='M-2222', cliente=cls.ana, marca='Yamaha', modelo='FZ', anio=2018, color='Negro',
            activo=False,
        )
        cls.url = reverse('usuarios:clientes_exportar')

    def _filas(self, respuesta):
        libro = load_workbook(BytesIO(respuesta.content))
        hoja = libro.active
        return [[celda for celda in fila] for fila in hoja.iter_rows(values_only=True)]

    def test_admin_descarga_un_xlsx_valido(self):
        self.client.force_login(self.admin)
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment; filename="clientes_vehiculos_', resp['Content-Disposition'])
        self.assertTrue(resp['Content-Disposition'].endswith('.xlsx"'))
        self.assertTrue(resp.content.startswith(b'PK'))

    def test_encabezados_incluyen_contacto_y_datos_basicos_del_vehiculo(self):
        self.client.force_login(self.admin)
        encabezados = self._filas(self.client.get(self.url))[0]
        for columna in ['DUI', 'Nombre', 'Apellido', 'Teléfono', 'Correo', 'Dirección',
                        'Placa', 'Marca', 'Modelo', 'Año', 'Color']:
            self.assertIn(columna, encabezados)

    def test_una_fila_por_vehiculo_y_cliente_sin_vehiculos_aparece_vacio(self):
        self.client.force_login(self.admin)
        filas = self._filas(self.client.get(self.url))[1:]
        # Ana (2 motos) + Beto (sin motos) = 3 filas; no incluye admin ni mecánico.
        self.assertEqual(len(filas), 3)
        placas = [f[7] for f in filas]
        self.assertEqual(placas, ['M-1111', 'M-2222', None])
        fila_ana = filas[0]
        self.assertEqual(fila_ana[:6], [
            '11111111-1', 'Ana', 'Álvarez', '7111-1111', 'ana@example.com', 'San Salvador',
        ])
        self.assertEqual(fila_ana[7:12], ['M-1111', 'Honda', 'CB190', 2020, 'Rojo'])
        fila_beto = filas[2]
        self.assertEqual(fila_beto[0], '22222222-2')
        self.assertEqual(fila_beto[7:], [None] * 6)

    def test_incluye_clientes_y_vehiculos_inactivos_con_su_estado(self):
        self.client.force_login(self.admin)
        filas = self._filas(self.client.get(self.url))[1:]
        self.assertEqual(filas[0][6], 'Activo')
        self.assertEqual(filas[0][12], 'Activo')
        self.assertEqual(filas[1][12], 'Inactivo')   # moto inactiva
        self.assertEqual(filas[2][6], 'Inactivo')    # cliente inactivo

    def test_la_exportacion_no_depende_de_la_busqueda(self):
        self.client.force_login(self.admin)
        resp = self.client.get(self.url, {'q': 'texto-que-no-existe'})
        self.assertEqual(len(self._filas(resp)) - 1, 3)

    def test_sin_clientes_solo_trae_encabezados(self):
        Motocicleta.objects.all().delete()
        Cliente.objects.all().delete()
        self.client.force_login(self.admin)
        filas = self._filas(self.client.get(self.url))
        self.assertEqual(len(filas), 1)

    def test_un_dato_que_empieza_con_igual_no_se_convierte_en_formula(self):
        self.ana.direccion = '=1+1'
        self.ana.save()
        self.client.force_login(self.admin)
        libro = load_workbook(BytesIO(self.client.get(self.url).content))
        celda = libro.active['F2']
        self.assertEqual(celda.value, '=1+1')
        self.assertEqual(celda.data_type, 's')

    def test_el_excel_no_trae_filtros_ni_tablas_dinamicas(self):
        self.client.force_login(self.admin)
        libro = load_workbook(BytesIO(self.client.get(self.url).content))
        hoja = libro.active
        self.assertIsNone(hoja.auto_filter.ref)
        self.assertEqual(len(hoja.tables), 0)

    def test_cliente_no_puede_exportar(self):
        self.client.force_login(self.ana)
        resp = self.client.get(self.url)
        self.assertRedirects(resp, reverse('core:home'), fetch_redirect_response=False)

    def test_mecanico_no_puede_exportar(self):
        self.client.force_login(self.mecanico)
        resp = self.client.get(self.url)
        self.assertRedirects(resp, reverse('core:home'), fetch_redirect_response=False)

    def test_anonimo_va_al_login(self):
        resp = self.client.get(self.url)
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse('usuarios:login'), resp['Location'])

    def test_boton_visible_en_la_seccion_de_clientes(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('usuarios:clientes_lista'))
        self.assertContains(resp, 'Exportar a Excel')
        self.assertContains(resp, self.url)

    def test_el_detalle_de_cliente_sigue_funcionando(self):
        self.client.force_login(self.admin)
        resp = self.client.get(reverse('usuarios:cliente_detalle', kwargs={'dui': '11111111-1'}))
        self.assertEqual(resp.status_code, 200)